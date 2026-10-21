"""HU-16 · Exportar la conciliación a Excel (export/excel.py)."""

from io import BytesIO

import pytest
from conciliasii.export.excel import generar_excel
from conciliasii.workflow import approval
from openpyxl import load_workbook

from .conftest import ANALISTA, GAF, JEFA, TESORERO, llevar_a_estado

HOJAS_DE_CATEGORIA = [
    "Cuadrado",
    "Diferencia de monto",
    "Diferencia de período",
    "Solo en SII",
    "Solo en ERP",
]


def exportar(repo, ejecucion) -> bytes:
    vigente = repo.ejecucion_vigente(ejecucion.periodo_id)
    return generar_excel(
        vigente,
        repo.listar_resultados(vigente.id),
        sugerencias=repo.listar_sugerencias(vigente.periodo_id),
        comentarios=repo.listar_comentarios(vigente.periodo_id, hasta_version=vigente.version),
        decisiones=repo.listar_decisiones(vigente.id),
    )


def libro(contenido: bytes):
    return load_workbook(BytesIO(contenido), data_only=True)


def filas(hoja) -> list[list]:
    return [list(f) for f in hoja.iter_rows(values_only=True)]


def textos(hoja) -> str:
    """Todo el texto visible de una hoja, para buscar valores sin depender de la disposición."""
    return " | ".join(str(c) for f in filas(hoja) for c in f if c is not None)


def test_descarga_un_xlsx_con_resumen_y_una_hoja_por_categoria(repo, ejecucion):
    contenido = exportar(repo, ejecucion)

    assert contenido[:2] == b"PK"  # un .xlsx es un zip
    nombres = libro(contenido).sheetnames
    assert nombres[0] == "Resumen"
    assert set(HOJAS_DE_CATEGORIA) <= set(nombres)


def test_la_hoja_resumen_tiene_cantidad_y_montos_por_categoria(repo, ejecucion):
    resumen = libro(exportar(repo, ejecucion))["Resumen"]

    por_categoria = {f[0]: f for f in filas(resumen) if f and f[0] in HOJAS_DE_CATEGORIA}
    assert set(por_categoria) == set(HOJAS_DE_CATEGORIA)
    assert por_categoria["Solo en SII"][1:4] == [4, 535_240, 0]
    assert por_categoria["Solo en ERP"][1:4] == [5, 0, 1_975_470]
    assert por_categoria["Cuadrado"][1:4] == [12, 4_702_313, 4_702_313]


def test_cada_hoja_de_categoria_tiene_una_fila_por_documento(repo, ejecucion):
    wb = libro(exportar(repo, ejecucion))

    esperado = {"Cuadrado": 12, "Diferencia de monto": 6, "Diferencia de período": 3}
    esperado |= {"Solo en SII": 4, "Solo en ERP": 5}
    for nombre, cantidad in esperado.items():
        hoja = wb[nombre]
        assert len(filas(hoja)) == cantidad + 1, nombre  # encabezado + documentos


def test_el_detalle_muestra_valores_de_ambos_lados_y_los_campos_distintos(repo, ejecucion):
    hoja = libro(exportar(repo, ejecucion))["Diferencia de monto"]
    encabezado, *cuerpo = filas(hoja)

    fila = next(f for f in cuerpo if "77300031-K|33|2211" in [str(c) for c in f])
    texto = " ".join(str(c) for c in fila)
    assert "95000" in texto and "90000" in texto  # IVA según SII y según BC
    assert "monto_iva_recuperable" in texto.lower() or "iva" in texto.lower()
    assert any("sii" in str(c).lower() for c in encabezado)
    assert any("erp" in str(c).lower() or "bc" in str(c).lower() for c in encabezado)


def test_el_excel_incluye_comentarios_con_autor_rol_fecha_y_hora(repo, ejecucion):
    approval.agregar_comentario(
        repo, ejecucion.id, "77300031-K|33|2211", TESORERO, "IVA mal digitado"
    )

    wb = libro(exportar(repo, ejecucion))

    detalle = textos(wb["Diferencia de monto"])
    assert "IVA mal digitado" in detalle
    comentarios = textos(wb["Comentarios"])
    assert TESORERO.nombre in comentarios
    assert "Tesorero" in comentarios
    assert "2026" in comentarios  # fecha y hora del comentario


def test_el_resumen_muestra_el_estado_y_quien_aprobo_y_cuando(repo, ejecucion):
    llevar_a_estado(repo, ejecucion, "Aprobado")

    resumen = textos(libro(exportar(repo, ejecucion))["Resumen"])

    assert "Aprobado" in resumen
    assert JEFA.nombre in resumen
    assert GAF.nombre in resumen


def test_el_resumen_muestra_el_estado_borrador_sin_aprobadores(repo, ejecucion):
    resumen = textos(libro(exportar(repo, ejecucion))["Resumen"])

    assert "Borrador" in resumen
    assert JEFA.nombre not in resumen


@pytest.mark.parametrize("actor", [TESORERO, ANALISTA, JEFA, GAF])
def test_los_cuatro_roles_de_negocio_pueden_descargar(actor):
    for estado in (
        "Borrador",
        "EnRevisionContable",
        "PendienteJefaContabilidad",
        "PendienteGAF",
        "Aprobado",
    ):
        assert "descargar_excel" in approval.acciones_permitidas(actor.rol, estado)


def test_sin_conciliacion_la_descarga_esta_deshabilitada():
    from conciliasii.domain.models import Rol

    for rol in (Rol.TESORERO, Rol.ANALISTA_CONTABLE, Rol.JEFA_CONTABILIDAD, Rol.GAF):
        assert "descargar_excel" not in approval.acciones_permitidas(rol, None)
